import csv

from openpyxl import load_workbook

from core.sra.exporters import HEADERS, export_csv, export_xlsx
from core.sra.generate import generate_risk_objectives
from core.sra.schemas import RiskObjectivesTable


def _table():
    return {
        "project_id": "p-1",
        "source_document_ids": ["doc.md"],
        "risk_objectives": [
            {
                "attacker_category": "Acteur étatique",
                "targeted_objectives": "Compromettre le segment sol pour perturber la mission.",
            },
            {
                "attacker_category": "Menace interne",
                "targeted_objectives": "Exfiltrer des données de télémétrie sensibles.",
            },
        ],
        "warnings": [],
        "generated_at": "2026-07-28T12:00:00",
    }


def test_generation_multi_documents_et_deduplication():
    answers = {
        "alpha": {
            "risk_objectives": [
                {
                    "categorie_attaquant": "Acteur étatique",
                    "extraits": ["alpha"],
                    "objectifs_vises": "Perturber le contrôle de la mission.",
                }
            ]
        },
        "beta": {
            "risk_objectives": [
                {
                    "categorie_attaquant": "Acteur étatique",
                    "extraits": ["beta"],
                    "objectifs_vises": "Perturber le contrôle de la mission.",
                },
                {
                    "categorie_attaquant": "Menace interne",
                    "extraits": ["beta"],
                    "objectifs_vises": "Exfiltrer les télémesures.",
                },
            ]
        },
    }

    def gather(document):
        return [{"content": document}]

    def llm(prompt, format=None):
        import json
        key = "alpha" if "alpha" in prompt else "beta"
        return json.dumps(answers[key])

    def pipeline(question, **kwargs):
        assert kwargs["prefilter"](question) == ["alpha", "beta"]
        assert kwargs["map_only"] is True
        extract = kwargs["extract"]
        extractions = [extract("alpha", question), extract("beta", question)]
        yield {"type": "map", "document": "alpha", "n_items": 1}
        yield {"type": "done", "result": {
            "extractions": extractions,
            "document_coverage": {"documents_without_evidence": []},
        }}

    def adjudicate(prompt, format=None):
        import json
        assert format == "json"
        payload = json.loads(prompt.split("[CANDIDATS]\n", 1)[1])
        return json.dumps({"decisions": [{
            "id": item["id"],
            "accepte": True,
            "categorie_normalisee": (
                "Acteur étatique"
                if item["categorie_extraite"] == "Acteur étatique"
                else "Menace interne malveillante"
            ),
            "objectifs_atomiques": [item["objectif_extrait"]],
            "motif_rejet": "",
            "confiance": "haute",
        } for item in payload]})

    table = generate_risk_objectives(
        "p-1", ["alpha", "beta"], pipeline=pipeline,
        extract_llm=llm, project_llm=adjudicate, gather=gather,
    )
    model = RiskObjectivesTable.model_validate(table)
    assert model.schema_version == "2.0"
    assert [row.ro_type_id for row in model.risk_origin_types] == ["RO-01", "RO-06"]
    assert len(model.risk_origin_types[0].evidence) == 2
    consolidation = next(e for e in model.trace if e["stage"] == "sra_consolidation")
    assert consolidation["extracted_candidates"] == 3
    assert consolidation["ro_types_retained"] == 2
    assert consolidation["occurrences_merged_into_types"] == 1


def test_adjudication_rejette_non_malveillant_et_atomise_objectifs():
    def pipeline(question, **kwargs):
        yield {"type": "done", "result": {
            "extractions": [{"candidates": [
                {
                    "attacker_category": "France",
                    "targeted_objectives": "Développer une capacité défensive.",
                    "evidence": "développe une capacité défensive",
                    "source": "doctrine.md",
                },
                {
                    "attacker_category": "Chine",
                    "targeted_objectives": "Espionner et perturber les communications.",
                    "evidence": "espionner et perturber les communications",
                    "source": "menace.md",
                },
            ]}],
            "document_coverage": {"documents_without_evidence": []},
        }}

    def adjudicate(prompt, format=None):
        import json
        return json.dumps({"decisions": [
            {
                "id": 1, "accepte": False, "categorie_normalisee": "",
                "objectifs_atomiques": [],
                "motif_rejet": "activité défensive légitime",
                "confiance": "haute",
            },
            {
                "id": 2, "accepte": True,
                "categorie_normalisee": "Acteur étatique (Chine)",
                "objectifs_atomiques": [
                    "Collecter clandestinement des informations de mission.",
                    "Priver la mission de ses communications.",
                ],
                "motif_rejet": "", "confiance": "haute",
            },
        ]})

    table = generate_risk_objectives(
        "p-2", ["doctrine.md", "menace.md"], pipeline=pipeline,
        project_llm=adjudicate,
    )
    assert len(table["risk_origin_types"]) == 1
    assert table["risk_origin_types"][0]["ro_type_id"] == "RO-01"
    assert table["risk_origin_types"][0]["observed_instances"] == ["Chine"]
    assert len(table["risk_origin_types"][0]["characterization"]) == 2
    event = next(
        item for item in table["trace"]
        if item["stage"] == "sra_adjudication_batch"
    )
    assert event["candidates_rejected"] == 1
    assert event["rejections"][0]["reason"] == "activité défensive légitime"


def test_adjudication_incomplete_rejette_par_securite():
    def pipeline(question, **kwargs):
        yield {"type": "done", "result": {
            "extractions": [{"candidates": [{
                "attacker_category": "Acteur",
                "targeted_objectives": "Objectif",
                "evidence": "preuve",
                "source": "source.md",
            }]}],
            "document_coverage": {"documents_without_evidence": []},
        }}

    table = generate_risk_objectives(
        "p-3", ["source.md"], pipeline=pipeline,
        project_llm=lambda prompt, format=None: '{"decisions": []}',
    )
    assert table["risk_objectives"] == []
    assert "rejeté(s) par sécurité" in table["warnings"][0]


def test_taxonomie_personnalisee_remplace_le_catalogue_par_defaut():
    def pipeline(question, **kwargs):
        yield {"type": "done", "result": {"extractions": [{"candidates": [{
            "attacker_category": "Opérateur hostile", "targeted_objectives": "Dégrader le service",
            "evidence": "dégrade le service", "source": "test-7.md",
        }]}], "document_coverage": {"documents_without_evidence": []}}}

    table = generate_risk_objectives(
        "p-7", ["test-7.md"], pipeline=pipeline,
        project_llm=lambda *_args, **_kwargs: '{"decisions":[{"id":1,"accepte":true,"categorie_normalisee":"Opérateur hostile","objectifs_atomiques":["Dégrader le service"],"motif_rejet":"","confiance":"haute"}]}',
        generation_config={"taxonomy_id": "test-7", "taxonomy_version": "7.0",
            "ro_types": [{"ro_type_id": "RO-77", "canonical_type": "Opérateur hostile",
                          "definition": "Acteur propre au septième corpus."}]},
    )
    assert [(row["ro_type_id"], row["canonical_type"]) for row in table["risk_origin_types"]] == [
        ("RO-77", "Opérateur hostile")]


def test_taxonomie_extensible_produit_un_identifiant_stable_a_revoir():
    def pipeline(question, **kwargs):
        yield {"type": "done", "result": {"extractions": [{"candidates": [{
            "attacker_category": "Collectif orbital autonome", "targeted_objectives": "Brouiller le signal",
            "evidence": "brouille le signal", "source": "test-7.md",
        }]}], "document_coverage": {"documents_without_evidence": []}}}
    adjudicate = lambda *_args, **_kwargs: '{"decisions":[{"id":1,"accepte":true,"categorie_normalisee":"Collectif orbital autonome","objectifs_atomiques":["Brouiller le signal"],"motif_rejet":"","confiance":"moyenne"}]}'
    config = {"taxonomy_mode": "extensible", "taxonomy_id": "test-7", "taxonomy_version": "7.0"}
    first = generate_risk_objectives("p-7", ["test-7.md"], pipeline=pipeline,
                                     project_llm=adjudicate, generation_config=config)
    second = generate_risk_objectives("p-8", ["test-7.md"], pipeline=pipeline,
                                      project_llm=adjudicate, generation_config=config)
    assert first["risk_origin_types"][0]["ro_type_id"] == second["risk_origin_types"][0]["ro_type_id"]
    assert first["risk_origin_types"][0]["ro_type_id"].startswith("RO-EXT-")
    assert "à valider" in first["risk_origin_types"][0]["review_comment"]


def test_export_csv_ro_types_six_colonnes(tmp_path):
    destination = tmp_path / "ro.csv"
    export_csv(_table(), destination)
    with destination.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert tuple(rows[0]) == HEADERS
    assert all(len(row) == 6 for row in rows)


def test_export_xlsx_ro_types_six_colonnes(tmp_path):
    destination = tmp_path / "ro.xlsx"
    export_xlsx(_table(), destination)
    sheet = load_workbook(destination).active
    assert sheet.title == "RO Types"
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet.max_row == 3
    assert sheet.max_column == 6


def test_extraction_ro_fractionne_seulement_si_sortie_tronquee():
    from core.llm_client import GenerationTruncated
    from core.sra.generate import extract_sra_from_document
    calls = []

    def llm(prompt, format=None):
        text = prompt.split("[TEXTE]\n", 1)[1]
        calls.append(len(text))
        if len(text) > 4_000:
            raise GenerationTruncated("sortie plafonnée")
        return "{\"risk_objectives\":[]}"

    result = extract_sra_from_document(
        "doc", "aspect", llm=llm,
        gather=lambda _: [{"content": "A" * 12_000}],
    )
    assert result["candidates"] == []
    assert calls[0] >= 12_000
    assert len(calls) > 1


def test_task_metrics_correlate_llm_calls_and_efficiency():
    from utils import task_metrics
    task = task_metrics.begin("chat", question="test")
    task_metrics.record_llm_call(task["task_id"], "model-a", "generate", 100, 25, 2.5)
    task_metrics.record_llm_call(task["task_id"], "model-b", "judge", 40, 5, 0.5)
    task_metrics.set_trajectory(task["task_id"], tool_calls=2, steps=3)
    result = task_metrics.finish(task["task_id"], completion="complete")
    assert result["trajectory"]["llm_calls"] == 2
    assert result["trajectory"]["llm_calls_by_role"] == {"generate": 1, "judge": 1}
    assert result["efficiency"]["total_tokens"] == 170
    assert result["quality"]["semantic_judge"] == "not_run"


def test_sra_task_metrics_detect_partial_document_coverage():
    from utils import task_metrics
    task = task_metrics.begin("sra.to")
    result = {"table": {"source_document_ids": ["a.pdf"], "target_objectives": [{"to_id": "TO-1"}], "warnings": []}}
    metrics = task_metrics.finish_sra(task["task_id"], result, ["a.pdf", "b.pdf"], "target_objectives")
    assert metrics["status"] == "partial"
    assert metrics["outcome"]["completion"] == "partial"
    assert metrics["outcome"]["processed_documents"] == 1
