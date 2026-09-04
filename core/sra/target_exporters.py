"""Exports de la table TO, consolidée à une ligne par RO Type."""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from core.sra.schemas import TargetObjectivesTable
TO_HEADERS = ("TO ID", "RO Type ID", "RO Type canonique", "Attaques", "Objectifs", "Cibles",
              "Propriétés de sécurité", "Domaines cibles", "Motivations", "Justifications", "Sources et preuves")
def _values(row):
    if hasattr(row, "attacks"):
        return (row.to_id, row.ro_type_id, row.ro_canonical_type,
            "\n".join(f"{item.scenario_id} — {item.attack}" for item in row.attacks),
            "\n".join(f"{item.scenario_id} — {item.objective}" for item in row.attacks),
            "\n".join(f"{item.scenario_id} — {', '.join(item.target_names)}" for item in row.attacks),
            "\n".join(f"{item.scenario_id} — {', '.join(item.security_properties)}" for item in row.attacks),
            "\n".join(f"{item.scenario_id} — {', '.join(item.target_domains)}" for item in row.attacks),
            "\n".join(f"{item.scenario_id} — {', '.join(item.motivations)}" for item in row.attacks),
            "\n".join(f"{item.scenario_id} — {item.rationale}" for item in row.attacks),
            "\n".join(f"{item.scenario_id} — {proof.source}: {proof.quote}"
                      for item in row.attacks for proof in item.evidence))
    return (row.to_id, row.ro_type_id, row.ro_canonical_type, "", "",
            row.target_name, "", "", "", row.rationale,
            "\n".join(f"{item.source}: {item.quote}" for item in row.evidence))
def export_target_csv(table: dict, destination: Path) -> None:
    model = TargetObjectivesTable.model_validate(table)
    with destination.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream); writer.writerow(TO_HEADERS)
        for row in model.target_objectives: writer.writerow(_values(row))
def export_target_xlsx(table: dict, destination: Path) -> None:
    model = TargetObjectivesTable.model_validate(table); workbook = Workbook(); sheet = workbook.active
    sheet.title = "TO par RO Type"; sheet.append(TO_HEADERS)
    for row in model.target_objectives: sheet.append(_values(row))
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip("ABCDEFGHIJK", (12, 12, 28, 42, 48, 55, 34, 34, 34, 80, 100)): sheet.column_dimensions[column].width = width
    for cell in sheet[1]: cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(destination)
CLASSIFIED_TO_HEADERS = (
    "TO décisionnel ID", "RO Type ID", "RO Type canonique", "Effet de sécurité",
    "Domaine principal", "Domaines secondaires",
    "Attaque générique", "Synthèse des attaques", "Synthèse des objectifs",
    "Scénarios sources", "Nombre de scénarios", "Cibles", "Propriétés de sécurité",
    "Motivations", "Sources et preuves",
)


def _classified_values(row, include_ro: bool = True):
    return (
        row.classification_id, row.ro_type_id if include_ro else "",
        row.ro_canonical_type if include_ro else "",
        row.effect_category, row.primary_target_domain,
        "\n".join(row.secondary_target_domains), row.generic_attack_name, row.attack_summary,
        row.objective_summary, "\n".join(row.scenario_ids),
        row.source_scenario_count, "\n".join(row.target_names),
        "\n".join(row.security_properties), "\n".join(row.motivations),
        "\n".join(f"{proof.source}: {proof.quote}" for proof in row.evidence),
    )


def export_classified_target_csv(table: dict, destination: Path) -> None:
    from core.sra.schemas import ClassifiedTargetObjectivesTable

    model = ClassifiedTargetObjectivesTable.model_validate(table)
    with destination.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(CLASSIFIED_TO_HEADERS)
        previous_ro_type_id = None
        for row in model.classified_objectives:
            writer.writerow(_classified_values(row, row.ro_type_id != previous_ro_type_id))
            previous_ro_type_id = row.ro_type_id


def export_classified_target_xlsx(table: dict, destination: Path) -> None:
    from core.sra.schemas import ClassifiedTargetObjectivesTable

    model = ClassifiedTargetObjectivesTable.model_validate(table)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TO décisionnels"
    sheet.append(CLASSIFIED_TO_HEADERS)
    group_start = 2
    previous_ro_type_id = None
    for row_index, row in enumerate(model.classified_objectives, 2):
        sheet.append(_classified_values(row))
        if previous_ro_type_id is not None and row.ro_type_id != previous_ro_type_id:
            if row_index - 1 > group_start:
                sheet.merge_cells(start_row=group_start, start_column=2, end_row=row_index - 1, end_column=2)
                sheet.merge_cells(start_row=group_start, start_column=3, end_row=row_index - 1, end_column=3)
            group_start = row_index
        previous_ro_type_id = row.ro_type_id
    final_row = len(model.classified_objectives) + 1
    if previous_ro_type_id is not None and final_row > group_start:
        sheet.merge_cells(start_row=group_start, start_column=2, end_row=final_row, end_column=2)
        sheet.merge_cells(start_row=group_start, start_column=3, end_row=final_row, end_column=3)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = (18, 12, 28, 30, 22, 35, 42, 70, 60, 36, 18, 55, 30, 30, 100)
    for index, width in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(destination)
